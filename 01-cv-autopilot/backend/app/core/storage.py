"""Couche : infrastructure stockage -- un seul endroit qui sait où vivent les
fichiers uploadés. Implémentation disque local, montée en production
(2026-08-15) sur un volume Railway persistant à `/app/storage` -- même
chemin qu'en local (`backend/storage/`), donc zéro changement de code pour
passer du dev à la prod. Testé en réel : fichier généré, redéploiement
forcé, fichier retéléchargé identique octet pour octet (voir DEPLOY.md).

**Choix documenté, pas un raccourci silencieux** : un volume Railway (1
disque, 1 instance) suffit pour une bêta privée à faible concurrence --
pas besoin d'un compte S3/R2 tiers pour ça. Limite connue de cette
approche : ne scale pas à plusieurs instances du service (le volume est
attaché à une seule). À remplacer par un stockage objet (S3/R2) avant de
scaler horizontalement -- sans changer la signature de ces fonctions,
storage_key resterait une clé opaque, juste résolue différemment.
Ne fait jamais : lire le contenu métier d'un CV (ça reste dans les Services
qui appellent extract_text)."""
from __future__ import annotations

import uuid
from pathlib import Path

STORAGE_ROOT = Path(__file__).resolve().parent.parent.parent / "storage" / "cv_uploads"
GENERATED_ROOT = Path(__file__).resolve().parent.parent.parent / "storage" / "cv_files"


def save_upload(user_id: uuid.UUID, file_name: str, content: bytes) -> str:
    """Enregistre le fichier, renvoie sa storage_key (chemin relatif à STORAGE_ROOT)."""
    user_dir = STORAGE_ROOT / str(user_id)
    user_dir.mkdir(parents=True, exist_ok=True)

    safe_name = f"{uuid.uuid4().hex[:8]}_{file_name}"
    (user_dir / safe_name).write_bytes(content)
    return f"{user_id}/{safe_name}"


def delete_upload(storage_key: str) -> None:
    """Supprime le fichier physique -- appelé quand l'utilisateur retire un
    CV de sa liste (voir onboarding_service.delete_cv_upload). N'échoue pas
    si le fichier est déjà absent (idempotent)."""
    path = STORAGE_ROOT / storage_key
    path.unlink(missing_ok=True)


def _iter_docx_block_items(parent):
    """Parcourt les paragraphes ET tableaux d'un document/cellule .docx dans
    leur ordre réel d'apparition -- `doc.paragraphs` seul ignore tout texte
    situé dans un tableau. Motif standard python-docx (non exposé par
    l'API publique de la lib)."""
    from docx.document import Document as _DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph

    parent_elm = parent.element.body if isinstance(parent, _DocxDocument) else parent._tc
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _extract_docx_text(doc) -> str:
    """Texte brut d'un .docx, tableaux inclus (2026-08-15) -- trouvé en réel :
    un CV utilisant une mise en page en tableau (très courant sur les
    modèles de CV gratuits, colonne compétences + colonne expérience) était
    extrait comme vide ("aucune expérience identifiable"), alors que tout le
    contenu utile vivait dans les cellules, jamais visibles depuis
    `doc.paragraphs` seul. Un niveau de récursion dans les cellules -- les
    tableaux imbriqués dans un tableau restent hors périmètre, cas rare pour
    un CV."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    lines: list[str] = []
    for block in _iter_docx_block_items(doc):
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                for cell in row.cells:
                    for sub_block in _iter_docx_block_items(cell):
                        if isinstance(sub_block, Paragraph) and sub_block.text.strip():
                            lines.append(sub_block.text.strip())
    return "\n".join(lines)


def extract_text(storage_key: str, mime_type: str) -> str:
    """Lit le fichier stocké et en extrait le texte brut. Formats supportés
    aujourd'hui : .docx et texte brut -- même périmètre que les prototypes
    (prototypes/cv-extraction/), pas de PDF pour l'instant."""
    path = STORAGE_ROOT / storage_key

    if mime_type in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or path.suffix.lower() == ".docx":
        import docx

        doc = docx.Document(str(path))
        return _extract_docx_text(doc)

    return path.read_text(encoding="utf-8")


def save_generated_docx(user_id: uuid.UUID, role_id: uuid.UUID, version: int, doc) -> str:
    """Enregistre un .docx généré (python-docx Document), renvoie sa storage_key."""
    role_dir = GENERATED_ROOT / str(user_id)
    role_dir.mkdir(parents=True, exist_ok=True)

    file_name = f"{role_id}_v{version}.docx"
    doc.save(str(role_dir / file_name))
    return f"{user_id}/{file_name}"


def read_generated_file(storage_key: str) -> bytes:
    return (GENERATED_ROOT / storage_key).read_bytes()


def generated_file_path(storage_key: str) -> Path:
    """Chemin réel du fichier -- pour un appelant qui a besoin d'un `Path`
    plutôt que des octets (ex: application_send_service, qui passe la pièce
    jointe à smtp_adapter sans jamais lire son contenu lui-même). storage_key
    reste une clé opaque pour l'appelant -- seul ce module sait la résoudre."""
    return GENERATED_ROOT / storage_key


def delete_generated_file(storage_key: str) -> None:
    """Utilisé à la suppression d'un rôle (2026-08-15). Idempotent."""
    (GENERATED_ROOT / storage_key).unlink(missing_ok=True)


def delete_user_storage(user_id: uuid.UUID) -> None:
    """Utilisé par la suppression admin d'un compte entier (2026-08-15,
    nettoyage des comptes de test en bêta). save_upload/save_generated_docx
    rangent tout sous un dossier par utilisateur -- supprimer ces deux
    dossiers couvre tous les CV uploadés et générés de ce compte, sans avoir
    à lister chaque storage_key un par un. Idempotent (pas d'erreur si le
    dossier n'existe pas ou est déjà vide)."""
    import shutil

    shutil.rmtree(STORAGE_ROOT / str(user_id), ignore_errors=True)
    shutil.rmtree(GENERATED_ROOT / str(user_id), ignore_errors=True)
