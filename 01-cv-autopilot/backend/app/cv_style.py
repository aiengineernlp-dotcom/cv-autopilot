"""Couche : Presentation.

Copie de `../engine/cv_style.py` -- délibérément dupliquée plutôt
qu'importée, pour que `backend/` reste déployable seul (Railway), sans
dépendre du dossier `engine/` de la Phase 1. Purement présentation
(tailles, couleurs, espacements) : aucune logique métier, donc peu de
risque de divergence -- si un changement visuel est fait ici, le reporter
dans `engine/cv_style.py` (et vice-versa).

Fait : primitives de construction d'un document Word (nom, titre, section, puce,
compétence...) — reproduit exactement le format mesuré dans les CV existants
(CV_Data_Analyst / Data_Engineer / Data_Scientist) : mêmes tailles, couleurs,
espacements. Générique : ne connaît aucun rôle, aucun domaine métier — un CV
juridique ou marketing utilise les mêmes fonctions qu'un CV ML/AI.
Ne fait jamais : savoir quel texte mettre (cv_generation_service), décider où
sauvegarder le fichier ou mettre à jour le statut d'un rôle (Service).
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

NAVY = RGBColor(0x1B, 0x4F, 0x72)
DARK = RGBColor(0x00, 0x00, 0x00)
BODY = RGBColor(0x2C, 0x2C, 0x2C)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT_GRAY = RGBColor(0x88, 0x88, 0x88)


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Inches(0.5903)
    section.right_margin = Inches(0.5903)
    section.top_margin = Inches(0.4514)
    section.bottom_margin = Inches(0.4514)
    return doc


def _run(paragraph, text, *, size, bold=False, color=BODY):
    r = paragraph.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = color
    return r


def add_name(doc, full_name: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1.75)
    _run(p, full_name, size=15, bold=True, color=NAVY)
    return p


def add_title(doc, title: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3.0)
    _run(p, title, size=11, bold=True, color=BODY)
    return p


def add_contact_line(doc, text: str, *, color, space_after: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    _run(p, text, size=7.5, color=color)
    return p


def add_section_header(doc, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8.0)
    p.paragraph_format.space_after = Pt(2.5)
    _run(p, title, size=11, bold=True, color=NAVY)
    return p


def add_body_paragraph(doc, text: str, *, space_after: float = 7.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    _run(p, text, size=8.5, color=BODY)
    return p


def add_entry_header(doc, title: str, sub: str, dates: str):
    """Ligne 'Poste | Entreprise | Dates' utilisée pour expérience/projet/éducation."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5.0)
    p.paragraph_format.space_after = Pt(2.5)
    _run(p, title, size=9.5, bold=True, color=DARK)
    _run(p, f"   |   {sub}", size=8.5, color=GRAY)
    _run(p, f"   |   {dates}", size=8.5, color=LIGHT_GRAY)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1.75)
    _run(p, f"• {text}", size=8.5, color=BODY)
    return p


def add_skill_line(doc, category: str, items: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2.0)
    _run(p, f"{category}: ", size=8.5, bold=True, color=NAVY)
    _run(p, items, size=8.5, color=BODY)
    return p
